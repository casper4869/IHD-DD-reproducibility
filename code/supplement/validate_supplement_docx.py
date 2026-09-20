#!/usr/bin/env python3
"""Validate the page-bounded supplementary tables against their source CSVs."""

from __future__ import annotations

import argparse
import importlib.util
from pathlib import Path

from docx import Document


def load_builder(path: Path):
    spec = importlib.util.spec_from_file_location("supplement_builder", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def extract_rows(tables, expected_columns):
    rows = []
    for table in tables:
        if len(table.columns) != expected_columns:
            raise ValueError(
                f"Expected {expected_columns} columns; found {len(table.columns)}"
            )
        rows.extend(tuple(cell.text.strip() for cell in row.cells) for row in table.rows[1:])
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", required=True, type=Path)
    parser.add_argument("--s1", required=True, type=Path)
    parser.add_argument("--s2", required=True, type=Path)
    parser.add_argument("--builder", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    args = parser.parse_args()

    builder = load_builder(args.builder)
    expected_s1 = builder.read_s1(args.s1)
    expected_s2 = builder.read_s2(args.s2)
    doc = Document(args.docx)
    if len(doc.tables) != 18:
        raise ValueError(f"Expected 18 page tables; found {len(doc.tables)}")
    actual_s1 = extract_rows(doc.tables[:9], 9)
    actual_s2 = extract_rows(doc.tables[9:], 6)

    checks = {
        "S1 row count is 204": len(actual_s1) == 204,
        "S2 row count is 204": len(actual_s2) == 204,
        "S1 cells match source CSV after display formatting": actual_s1 == expected_s1,
        "S2 cells match source CSV after display formatting": actual_s2 == expected_s2,
        "Document contains no inline shapes": len(doc.inline_shapes) == 0,
    }
    failures = [name for name, passed in checks.items() if not passed]
    lines = ["# Supplementary DOCX validation", ""]
    for name, passed in checks.items():
        lines.append(f"- {'PASS' if passed else 'FAIL'}: {name}")
    lines.extend(
        [
            "",
            f"- Tables: {len(doc.tables)}",
            f"- S1 displayed data rows: {len(actual_s1)}",
            f"- S2 displayed data rows: {len(actual_s2)}",
            "- Layout: one fixed-grid table per rendered page, 23 data rows per page except the final page in each table.",
        ]
    )
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text("\n".join(lines) + "\n", encoding="utf-8")
    if failures:
        raise SystemExit("Validation failed: " + "; ".join(failures))


if __name__ == "__main__":
    main()
