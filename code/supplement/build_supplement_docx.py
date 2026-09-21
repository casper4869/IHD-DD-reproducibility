#!/usr/bin/env python3
"""Build the revised supplementary tables as page-bounded Word tables.

The submitted supplement used two 205-row Word tables.  Word occasionally
collapsed the column grid on a continuation page during PDF conversion.  This
builder writes each visible page as a separate table with the same fixed
column grid, so every continuation page is deterministic and self-contained.
"""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


S1_HEADERS = (
    "Location",
    "Sex",
    "Pearson r",
    "P",
    "FDR q",
    "Sex",
    "Pearson r",
    "P",
    "FDR q",
)
S1_WIDTHS = (3.10, 0.80, 0.75, 0.65, 0.75, 0.80, 0.75, 0.65, 0.75)
S2_HEADERS = (
    "Location",
    "IHD→DD P",
    "DD→IHD P",
    "IHD→DD FDR q",
    "DD→IHD FDR q",
    "Temporal predictive classification",
)
S2_WIDTHS = (3.20, 1.00, 1.00, 1.15, 1.15, 1.90)


def format_p(value: str | float) -> str:
    number = float(value)
    return "<0.001" if number < 0.001 else f"{number:.3f}"


def set_cell_margins(cell, top=35, start=45, bottom=35, end=45) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 7.7) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def add_page_table(doc: Document, title: str, headers, rows, widths) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        cell.width = Inches(width)
        set_cell_text(cell, header, bold=True, size=7.5)
        shade_cell(cell, "D9EAF7")
    prevent_row_split(table.rows[0])

    for values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for index, (value, width) in enumerate(zip(values, widths)):
            cell = row.cells[index]
            cell.width = Inches(width)
            set_cell_text(cell, value)


def add_note(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(8)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def chunks(items, size):
    for start in range(0, len(items), size):
        yield items[start : start + size]


def read_s1(path: Path):
    grouped = defaultdict(dict)
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            grouped[row["location_name"]][row["sex_name"]] = row
    if len(grouped) != 204:
        raise ValueError(f"Expected 204 S1 locations; found {len(grouped)}")
    output = []
    for location in sorted(grouped):
        female = grouped[location].get("Female")
        male = grouped[location].get("Male")
        if female is None or male is None:
            raise ValueError(f"Missing Female or Male result for {location}")
        output.append(
            (
                location,
                "Female",
                f"{float(female['correlation']):.3f}",
                format_p(female["p_value"]),
                format_p(female["fdr_pvalue"]),
                "Male",
                f"{float(male['correlation']):.3f}",
                format_p(male["p_value"]),
                format_p(male["fdr_pvalue"]),
            )
        )
    return output


def read_s2(path: Path):
    output = []
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            raw_classification = row.get("classification") or row.get("direction") or ""
            classification = {
                "DD_to_IHD": "DD→IHD",
                "IHD_to_DD": "IHD→DD",
                "Neither": "No evidence",
                "NA": "No evidence",
            }.get(raw_classification, raw_classification)
            output.append(
                (
                    row["location_name"],
                    format_p(row.get("p_ihd_to_dd") or row["IHD_to_DD_p"]),
                    format_p(row.get("p_dd_to_ihd") or row["DD_to_IHD_p"]),
                    format_p(row.get("q_ihd_to_dd") or row["IHD_to_DD_q"]),
                    format_p(row.get("q_dd_to_ihd") or row["DD_to_IHD_q"]),
                    classification,
                )
            )
    if len(output) != 204:
        raise ValueError(f"Expected 204 S2 locations; found {len(output)}")
    return sorted(output)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.48)
    section.right_margin = Inches(0.48)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(8)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def build(s1_path: Path, s2_path: Path, output_path: Path, rows_per_page: int) -> None:
    s1_rows = read_s1(s1_path)
    s2_rows = read_s2(s2_path)
    s1_pages = list(chunks(s1_rows, rows_per_page))
    s2_pages = list(chunks(s2_rows, rows_per_page))

    doc = Document()
    configure_document(doc)
    s1_title = (
        "Table S1. Sex-stratified country-level correlations between age-specific incidence rates "
        "of ischaemic heart disease (IHD) and depressive disorders (DD) in 2021, with raw and "
        "FDR-adjusted P values."
    )
    for index, rows in enumerate(s1_pages):
        title = s1_title if index == 0 else "Table S1 (continued)."
        add_page_table(doc, title, S1_HEADERS, rows, S1_WIDTHS)
        if index == len(s1_pages) - 1:
            add_note(
                doc,
                "*Pearson r was calculated within each country and sex across 20 paired age-specific "
                "incidence rates in 2021. The 204 two-sided P values were adjusted separately within "
                "each sex using the Benjamini–Hochberg procedure applied to unrounded P values. "
                "FDR q<0.05 was considered statistically significant.",
            )
        page_break(doc)

    s2_title = (
        "Table S2. Country-level Granger temporal predictive precedence between IHD and DD, "
        "1992–2021: raw and FDR-adjusted P values for IHD→DD and DD→IHD."
    )
    for index, rows in enumerate(s2_pages):
        title = s2_title if index == 0 else "Table S2 (continued)."
        add_page_table(doc, title, S2_HEADERS, rows, S2_WIDTHS)
        if index == len(s2_pages) - 1:
            add_note(
                doc,
                "*Models used annual log changes, a common Schwarz-BIC lag of 1–3 per country, SDI "
                "change as an exogenous covariate, and finite-sample-adjusted Newey–West covariance. "
                "For each direction, the null hypothesis was that the lagged incidence terms of the "
                "putative predictor were jointly zero. The 204 P values were adjusted separately for "
                "the IHD→DD and DD→IHD families using the Benjamini–Hochberg procedure applied to "
                "unrounded values; classification used q<0.05. An arrow denotes temporal predictive "
                "precedence and does not imply causation. Fixed-lag, HC3, residual, and influence "
                "diagnostics are provided in the reproducibility package.",
            )
        if index != len(s2_pages) - 1:
            page_break(doc)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--s1", required=True, type=Path)
    parser.add_argument("--s2", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--rows-per-page", type=int, default=23)
    args = parser.parse_args()
    build(args.s1, args.s2, args.output, args.rows_per_page)


if __name__ == "__main__":
    main()
